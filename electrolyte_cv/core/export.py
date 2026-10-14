from __future__ import annotations

from pathlib import Path
import pandas as pd

from .config import EN_HEADERS, CN_HEADERS
from .logging_utils import log_warning

def _column_display_name(column):
    return f"{EN_HEADERS.get(column, column)}\n{CN_HEADERS.get(column, column)}"


def _style_summary_workbook(xlsx_path):
    """四工作表统一格式：双层标题、冻结前两行、居中、合理列宽与数字格式。"""
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        log_warning('未安装 openpyxl，主结果工作簿将保留基础格式')
        return

    wb = load_workbook(xlsx_path)
    thin = Side(style='thin', color='D9DEE7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    en_fill = PatternFill('solid', fgColor='1F4E78')
    cn_fill = PatternFill('solid', fgColor='D9EAF7')
    alt_fill = PatternFill('solid', fgColor='F7FAFC')

    for ws in wb.worksheets:
        ws.freeze_panes = 'A3'
        ws.sheet_view.showGridLines = False
        max_col = max(ws.max_column, 1)
        max_row = max(ws.max_row, 2)

        for cell in ws[1]:
            cell.fill = en_fill
            cell.font = Font(bold=True, color='FFFFFF', size=10.5)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        for cell in ws[2]:
            cell.fill = cn_fill
            cell.font = Font(bold=True, color='1F1F1F', size=10)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 30

        for r in range(3, max_row + 1):
            ws.row_dimensions[r].height = 22
            for c in range(1, max_col + 1):
                cell = ws.cell(r, c)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
                cell.border = border
                if r % 2 == 0:
                    cell.fill = alt_fill
                if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                    en_header = str(ws.cell(1, c).value or '')
                    if '(A)' in en_header or '(C)' in en_header or '(A·V)' in en_header:
                        cell.number_format = '0.000E+00'
                    elif '(mV)' in en_header:
                        cell.number_format = '+0.0;-0.0;0.0'
                    elif 'R2' in en_header or 'R²' in en_header or 'fraction' in en_header:
                        cell.number_format = '0.000'
                    elif '(V)' in en_header:
                        cell.number_format = '0.0000'

        for c in range(1, max_col + 1):
            en_header = str(ws.cell(1, c).value or '')
            cn_header = str(ws.cell(2, c).value or '')
            if en_header == 'Curve':
                width = 16
            elif 'method' in en_header.lower():
                width = 24
            elif 'status' in en_header.lower() or en_header.endswith('_QC') or en_header == 'Scan_QC':
                width = 18
            elif 'Reference' in en_header:
                width = 18
            elif 'direction_fraction' in en_header:
                width = 22
            elif 'difference' in en_header.lower() or 'shift' in en_header.lower():
                width = 24
            elif 'Noise_sigma' in en_header:
                width = 20
            elif 'window' in en_header.lower():
                width = 20
            elif '(V)' in en_header:
                width = 23
            elif '(A)' in en_header or '(C)' in en_header or '(A·V)' in en_header:
                width = 20
            else:
                width = min(max(max(len(en_header), len(cn_header)) + 3, 14), 26)
            ws.column_dimensions[get_column_letter(c)].width = width

    wb.save(xlsx_path)


def _write_main_results_workbook(xlsx_path, df, mode):
    """主结果工作簿固定为四个工作表；Onset 两种方法独立输出。"""
    if mode == 'max':
        peak_col, current_col = 'E_ox_peak (V)', 'I_ox_peak (A)'
        onset_threshold_col, onset_tangent_col = 'E_onset_ox_threshold (V)', 'E_onset_ox_tangent (V)'
        peak_shift_col = 'Ox_peak_shift_vs_ref (mV)'
        onset_threshold_shift_col = 'Ox_onset_threshold_shift_vs_ref (mV)'
        onset_tangent_shift_col = 'Ox_onset_tangent_shift_vs_ref (mV)'
        charge_col = 'Q_ox (C)'
    else:
        peak_col, current_col = 'E_red_peak (V)', 'I_red_peak (A)'
        onset_threshold_col, onset_tangent_col = 'E_onset_red_threshold (V)', 'E_onset_red_tangent (V)'
        peak_shift_col = 'Red_peak_shift_vs_ref (mV)'
        onset_threshold_shift_col = 'Red_onset_threshold_shift_vs_ref (mV)'
        onset_tangent_shift_col = 'Red_onset_tangent_shift_vs_ref (mV)'
        charge_col = 'Q_red (C)'

    sheets = [
        ('01_Charge_Integral', [
            'Curve', 'Q1 (C)', 'Q2 (C)', charge_col,
            'A1 (A·V)', 'A2 (A·V)', 'A_total (A·V)',
            'V_start_Part1', 'V_end_Part1', 'V_start_Part2', 'V_end_Part2'
        ]),
        ('02_Potential_Features', [
            'Curve', peak_col, current_col, onset_threshold_col, onset_tangent_col
        ]),
        ('03_Potential_Shift', [
            'Curve', 'Reference_curve', peak_shift_col,
            onset_threshold_shift_col, onset_tangent_shift_col
        ]),
        ('04_Method_QC', [
            'Curve', 'Peak_status', 'Threshold_QC', 'Tangent_QC', 'Scan_QC',
            'Scan_direction_fraction', 'Onset_difference (mV)', 'Onset_tangent_R2',
            'Baseline_R2', 'Noise_sigma (A)', 'Smooth_window_points'
        ]),
    ]

    try:
        with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
            for sheet_name, columns in sheets:
                existing = [c for c in columns if c in df.columns]
                sub = df[existing].copy()
                sub.to_excel(writer, sheet_name=sheet_name, index=False, header=False, startrow=2)
                ws = writer.book[sheet_name]
                for col_idx, column in enumerate(existing, 1):
                    ws.cell(1, col_idx, EN_HEADERS.get(column, column))
                    ws.cell(2, col_idx, CN_HEADERS.get(column, column))
    except PermissionError as exc:
        raise PermissionError(f'结果文件保存失败：{xlsx_path}') from exc
    _style_summary_workbook(xlsx_path)


def _style_simple_data_workbook(xlsx_path):
    """对原始/验证类工作簿做轻量格式化，避免影响大文件处理速度。"""
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return
    try:
        wb = load_workbook(xlsx_path)
        for ws in wb.worksheets:
            ws.freeze_panes = 'A2'
            ws.sheet_view.showGridLines = False
            if ws.max_row >= 1:
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill('solid', fgColor='D9EAF7')
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                ws.row_dimensions[1].height = 26
            # CV原始表通常只有2列，限制为前12列防止异常文件造成格式化过慢。
            style_cols = min(ws.max_column, 12)
            for c in range(1, style_cols + 1):
                ws.column_dimensions[chr(64 + c) if c <= 26 else 'A'].width = 16
            # 所有已使用单元格居中。20圈数据量仍可接受。
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=style_cols):
                for cell in row:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
        wb.save(xlsx_path)
    except Exception as exc:
        log_warning(f'工作簿格式化跳过：{Path(xlsx_path).name}，{exc}')
